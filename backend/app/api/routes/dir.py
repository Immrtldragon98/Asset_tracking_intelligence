import io
from datetime import date
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel, Field
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.database.session import get_db
from app.models.drive_inventory import DriveInventory

router = APIRouter()


def _text(value: Any):
    if value is None:
        return None
    result = str(value).strip()
    return result or None


def _number(value: Any):
    if value is None or str(value).strip() == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _serialize(row: DriveInventory):
    return {
        "id": row.id,
        "source_row": row.source_row,
        "asset_code": f"DIR-{row.source_row:03d}",
        "name": row.equipment_name,
        "area": row.area,
        "location": row.location,
        "supply_from": row.supply_from,
        "panel_location": row.panel_location,
        "manufacturer": row.vfd_make,
        "rating": row.vfd_rating,
        "module_quantity": row.module_quantity,
        "model": row.model_no,
        "spare_available": row.spare_available,
        "spare_location": row.spare_location,
        "ip_address": row.ip_address,
        "module_power_supply": row.module_power_supply,
        "motor_kw": row.motor_kw,
        "voltage": row.voltage,
        "motor_rpm": row.motor_rpm,
        "motor_flc_amp": row.motor_flc_amp,
        "status": "REGISTERED",
        "criticality": "UNASSESSED",
        "notes": row.notes,
        "installation_date": row.installation_date.isoformat() if row.installation_date else None,
        "running_hours": row.running_hours or 0,
        "operational_status": row.operational_status or "RUNNING",
    }


@router.get("/drives")
def list_drives(db: Session = Depends(get_db)):
    rows = db.query(DriveInventory).order_by(DriveInventory.source_row).all()
    makes = sorted({r.vfd_make for r in rows if r.vfd_make})
    areas = sorted({r.area for r in rows if r.area})
    return {
        "module": {"code": "DIR", "name": "Drive Intelligent Register", "area": "CH#2"},
        "total": len(rows),
        "distinct_makes": len(makes),
        "distinct_areas": len(areas),
        "assets": [_serialize(row) for row in rows],
    }


@router.post("/import")
async def import_workbook(file: UploadFile = File(...), db: Session = Depends(get_db)):
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Upload an .xlsx workbook.")
    content = await file.read()
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Could not read the Excel workbook.") from exc

    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip().upper() for value in next(rows, ())]
    aliases = {
        "area": ("AREA",),
        "location": ("LOCATION",),
        "equipment_name": ("EQUIPMENT NAME",),
        "supply_from": ("SUPPLY FROM",),
        "panel_location": ("PANEL LOCATION",),
        "vfd_make": ("VFD MAKE",),
        "vfd_rating": ("VFD REATING", "VFD RATING"),
        "module_quantity": ("MODULE USED QUINTITY", "MODULE USED QUANTITY"),
        "model_no": ("MODEL NO.", "MODEL NO"),
        "spare_available": ("SPARE AVAILABULE", "SPARE AVAILABLE"),
        "spare_location": ("LOCATION.1", "SPARE LOCATION"),
        "ip_address": ("IP ADRESS", "IP ADDRESS"),
        "module_power_supply": ("VFD MODULE I/C POWER SUPPLY",),
        "motor_kw": ("MOTOR KW",),
        "voltage": ("VOLTAGE",),
        "motor_rpm": ("MOTOR RPM",),
        "motor_flc_amp": ("MOTOR FLC(AMP)", "MOTOR FLC (AMP)"),
    }
    indexes = {}
    for key, names in aliases.items():
        for name in names:
            if name in headers:
                indexes[key] = headers.index(name)
                break
    if "equipment_name" not in indexes:
        workbook.close()
        raise HTTPException(status_code=400, detail="The first worksheet does not contain an EQUIPMENT NAME column.")

    imported = 0
    skipped = 0
    try:
        for row_number, values in enumerate(rows, start=2):
            if not values or all(v is None or str(v).strip() == "" for v in values):
                continue
            data = {key: (values[idx] if idx < len(values) else None) for key, idx in indexes.items()}
            equipment = _text(data.get("equipment_name"))
            if not equipment:
                skipped += 1
                continue
            source_row = row_number - 1
            record = db.query(DriveInventory).filter(DriveInventory.source_row == source_row).one_or_none()
            if record is None:
                record = DriveInventory(source_row=source_row, equipment_name=equipment)
                db.add(record)
            record.equipment_name = equipment
            for field in ("area", "location", "supply_from", "panel_location", "vfd_make", "vfd_rating",
                          "module_quantity", "model_no", "spare_available", "spare_location", "ip_address",
                          "module_power_supply", "voltage"):
                setattr(record, field, _text(data.get(field)))
            for field in ("motor_kw", "motor_rpm", "motor_flc_amp"):
                setattr(record, field, _number(data.get(field)))
            imported += 1
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        workbook.close()
    return {"ok": True, "imported": imported, "skipped": skipped, "total": db.query(DriveInventory).count()}


class DriveUsageUpdate(BaseModel):
    installation_date: date | None = None
    running_hours: float = Field(ge=0)


@router.patch("/drives/{drive_id}/usage")
def update_drive_usage(drive_id: int, payload: DriveUsageUpdate, db: Session = Depends(get_db)):
    row = db.query(DriveInventory).filter(DriveInventory.id == drive_id).one_or_none()
    if row is None:
        raise HTTPException(status_code=404, detail="Drive record not found.")
    row.installation_date = payload.installation_date
    row.running_hours = payload.running_hours
    db.commit()
    db.refresh(row)
    return _serialize(row)


class DriveStatusUpdate(BaseModel):
    operational_status: str


@router.patch("/drives/{drive_id}/status")
def update_drive_status(drive_id: int, payload: DriveStatusUpdate, db: Session = Depends(get_db)):
    row = db.query(DriveInventory).filter(DriveInventory.id == drive_id).one_or_none()
    if row is None:
        raise HTTPException(status_code=404, detail="Drive record not found.")
    status = payload.operational_status.strip().upper()
    if status not in {"RUNNING", "STOPPED", "MAINTENANCE", "FAULT"}:
        raise HTTPException(status_code=400, detail="Status must be RUNNING, STOPPED, MAINTENANCE, or FAULT.")
    row.operational_status = status
    db.commit()
    db.refresh(row)
    return _serialize(row)
